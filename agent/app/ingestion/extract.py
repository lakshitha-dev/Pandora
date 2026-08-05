"""Text extraction and ingestion for user-uploaded documents.

The brief requires the application to accept uploads in at least one common
format; we support PDF, DOCX, CSV, and TXT/MD.

Chunking mode is chosen by inspection, not assumption. An uploaded document
is scanned for the corpus's record-ID pattern: if it carries 50 or more, it
gets the record-aware chunker; otherwise it falls back to narrative
chunking. That threshold is in docs/API_CONTRACT.md §2.1 — an arbitrary
uploaded field report is narrative, the Pandora corpus is not.
"""

from __future__ import annotations

import base64
import csv
import io
import logging
import re
from dataclasses import dataclass

from app.ingestion.parse import (
    Chunk,
    _pack,
    _slug,
    classify_evidence_quality,
)

logger = logging.getLogger(__name__)

RECORD_PATTERN = re.compile(r"\b(?:REG|STL|FAU|FLR|MED|ACC|INC|POL|KC)-\d+\b")
RECORD_AWARE_THRESHOLD = 50

SUPPORTED_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/csv": "csv",
    "text/plain": "txt",
    "text/markdown": "txt",
}
EXTENSION_TYPES = {
    ".pdf": "pdf", ".docx": "docx", ".csv": "csv", ".txt": "txt", ".md": "txt",
}


class UnsupportedFileType(ValueError):
    """Raised for a file type we cannot extract text from."""


@dataclass
class ExtractedDocument:
    text: str
    page_count: int | None
    kind: str


def _detect_kind(file_name: str, content_type: str) -> str:
    if content_type in SUPPORTED_TYPES:
        return SUPPORTED_TYPES[content_type]
    lowered = file_name.lower()
    for ext, kind in EXTENSION_TYPES.items():
        if lowered.endswith(ext):
            return kind
    raise UnsupportedFileType(
        f"unsupported file type {content_type!r} for {file_name!r}; "
        "supported: PDF, DOCX, CSV, TXT/MD"
    )


# ─────────────────────────────────────────────────────────────────────────
# Per-format extraction
# ─────────────────────────────────────────────────────────────────────────


def _extract_pdf(data: bytes) -> ExtractedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - one bad page must not fail the doc
            pages.append("")

    # Page markers survive into the chunker so citations can carry a page.
    text = "\n\n".join(f"[[PAGE {i + 1}]]\n{t}" for i, t in enumerate(pages) if t.strip())
    return ExtractedDocument(text=text, page_count=len(reader.pages), kind="pdf")


def _extract_docx(data: bytes) -> ExtractedDocument:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        # Map Word heading styles onto markdown so heading-aware splitting works.
        style = (p.style.name or "").lower() if p.style else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            parts.append(f"{'#' * min(int(level), 3)} {p.text.strip()}")
        else:
            parts.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" · ".join(cells))

    return ExtractedDocument(text="\n\n".join(parts), page_count=None, kind="docx")


def _extract_csv(data: bytes) -> ExtractedDocument:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return ExtractedDocument(text="", page_count=None, kind="csv")

    header, *body = rows
    # One line per row, labelled with its column names — a row stays
    # meaningful on its own once retrieved out of context.
    lines = [
        " · ".join(f"{h.strip()}: {v.strip()}" for h, v in zip(header, row) if v.strip())
        for row in body
        if any(v.strip() for v in row)
    ]
    return ExtractedDocument(text="\n\n".join(lines), page_count=None, kind="csv")


def _extract_txt(data: bytes) -> ExtractedDocument:
    return ExtractedDocument(
        text=data.decode("utf-8", errors="replace"), page_count=None, kind="txt"
    )


def extract_text(content_base64: str, file_name: str, content_type: str) -> ExtractedDocument:
    """Decode an uploaded file and extract its text."""
    kind = _detect_kind(file_name, content_type)
    try:
        data = base64.b64decode(content_base64)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"content_base64 is not valid base64: {exc}") from exc

    extractor = {
        "pdf": _extract_pdf, "docx": _extract_docx,
        "csv": _extract_csv, "txt": _extract_txt,
    }[kind]
    doc = extractor(data)

    if not doc.text.strip():
        raise ValueError(
            f"no text extracted from {file_name!r}. Scanned/image-only PDFs are "
            "not supported (no OCR)."
        )
    return doc


# ─────────────────────────────────────────────────────────────────────────
# Chunking uploaded text
# ─────────────────────────────────────────────────────────────────────────

_PAGE_MARKER = re.compile(r"\[\[PAGE (\d+)\]\]")
_HEADING = re.compile(r"^(#{1,3})\s+(.+)$", re.MULTILINE)


def chunk_uploaded(
    text: str, document_id: str, document_name: str
) -> tuple[list[Chunk], str, int]:
    """Chunk an uploaded document.

    Returns (chunks, chunking_mode, record_count).
    """
    record_count = len(set(RECORD_PATTERN.findall(text)))

    if record_count >= RECORD_AWARE_THRESHOLD:
        # The upload is corpus-shaped; reuse the record-aware path.
        from app.ingestion.parse import parse_corpus
        import tempfile
        from pathlib import Path

        with tempfile.NamedTemporaryFile(
            "w", suffix=".md", delete=False, encoding="utf-8"
        ) as fh:
            fh.write(text)
            tmp = Path(fh.name)
        try:
            chunks = parse_corpus(tmp)
        finally:
            tmp.unlink(missing_ok=True)

        for c in chunks:
            c.document_id = document_id
            c.document_name = document_name
            c.chunk_id = f"{document_id[:8]}-{c.chunk_id}"
        return chunks, "record_aware", record_count

    # Narrative: split on headings, then pack to the token budget.
    chunks: list[Chunk] = []
    page_by_offset: list[tuple[int, int]] = [
        (m.start(), int(m.group(1))) for m in _PAGE_MARKER.finditer(text)
    ]

    def page_at(offset: int) -> int | None:
        page = None
        for pos, num in page_by_offset:
            if pos <= offset:
                page = num
            else:
                break
        return page

    clean = _PAGE_MARKER.sub("", text)
    boundaries = [m.start() for m in _HEADING.finditer(clean)]
    boundaries = sorted({0, *boundaries, len(clean)})

    idx = 0
    for i in range(len(boundaries) - 1):
        piece = clean[boundaries[i] : boundaries[i + 1]].strip()
        if len(piece) < 120:
            continue
        first = piece.splitlines()[0].lstrip("# ").strip()
        heading = first[:80] if len(first) < 120 else ""

        for window in _pack(piece):
            chunks.append(
                Chunk(
                    chunk_id=f"{document_id[:8]}-up-{idx}-{_slug(heading or 'section')}",
                    content=(f"[{document_name}] {heading}\n\n{window}" if heading else window),
                    record_id=None,
                    record_type="narrative",
                    title=heading or document_name,
                    chapter=document_name,
                    section=heading,
                    evidence_quality=classify_evidence_quality("narrative", window),
                    document_id=document_id,
                    document_name=document_name,
                    page=page_at(boundaries[i]),
                )
            )
            idx += 1

    return chunks, "narrative", record_count
